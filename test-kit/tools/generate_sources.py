"""Generate the PDF and DOCX test sources from their Markdown references.

Development tool only: never called by build.bat. The generated files are
committed next to their Markdown source so that the build needs no dependency.

Usage (from the repository root, after creating the venv described in
test-kit/tools/README.md):

    test-kit/tools/.venv/Scripts/python test-kit/tools/generate_sources.py

The Markdown subset handled here is deliberately small: ATX headings, paragraphs,
bullet and numbered lists, pipe tables, horizontal rules, blank lines, and the
inline marks **bold** and *italic*. It is enough for fictional resumes, job
postings and certificates; it is not a general Markdown converter.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (
    HRFlowable,
    ListFlowable,
    ListItem,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

KIT_ROOT = Path(__file__).resolve().parents[1]
SOURCES = KIT_ROOT / "sources"

# Markdown reference -> output format. TXT sources are used as they are.
MANIFEST: dict[str, str] = {
    "profile/cv-nadia-berkani.md": "pdf",
    "profile/linkedin-nadia-berkani.md": "pdf",
    "profile/certification-cloud-platform-associate.md": "pdf",
    "profile/notes-complementaires-carriere.md": "docx",
    "profile/livret-formation-architecture-systemes-distribues.md": "pdf",
    "opportunities/001-lumen-pay-offre-developpeuse-backend-senior.md": "pdf",
    "opportunities/002-northwind-ledger-senior-software-engineer.md": "pdf",
}


# --------------------------------------------------------------------------- #
# Minimal Markdown block parser
# --------------------------------------------------------------------------- #


@dataclass
class Block:
    kind: str  # heading | paragraph | bullets | numbers | table | rule
    text: str = ""
    level: int = 0
    items: list[str] = field(default_factory=list)
    rows: list[list[str]] = field(default_factory=list)


def parse_markdown(text: str) -> list[Block]:
    blocks: list[Block] = []
    lines = text.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        if not stripped:
            i += 1
            continue
        if stripped == "---":
            blocks.append(Block("rule"))
            i += 1
            continue
        m = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if m:
            blocks.append(Block("heading", m.group(2), level=len(m.group(1))))
            i += 1
            continue
        if stripped.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{3,}:?", c) for c in cells):
                    rows.append(cells)
                i += 1
            blocks.append(Block("table", rows=rows))
            continue
        if re.match(r"^[-*]\s+", stripped):
            items = []
            while i < len(lines) and re.match(r"^\s*[-*]\s+", lines[i]):
                items.append(re.sub(r"^\s*[-*]\s+", "", lines[i]))
                i += 1
            blocks.append(Block("bullets", items=items))
            continue
        if re.match(r"^\d+\.\s+", stripped):
            items = []
            while i < len(lines) and re.match(r"^\s*\d+\.\s+", lines[i]):
                items.append(re.sub(r"^\s*\d+\.\s+", "", lines[i]))
                i += 1
            blocks.append(Block("numbers", items=items))
            continue
        # Paragraph: consecutive non-empty lines that start no other block.
        para = []
        while i < len(lines) and lines[i].strip() and not _starts_block(lines[i]):
            para.append(lines[i].strip())
            i += 1
        blocks.append(Block("paragraph", "\n".join(para)))
    return blocks


def _starts_block(line: str) -> bool:
    s = line.strip()
    return (
        s == "---"
        or bool(re.match(r"^#{1,6}\s+", s))
        or s.startswith("|")
        or bool(re.match(r"^[-*]\s+", s))
        or bool(re.match(r"^\d+\.\s+", s))
    )


# Inline marks: returns a list of (text, bold, italic) runs.
_INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*)")


def inline_runs(text: str) -> list[tuple[str, bool, bool]]:
    runs = []
    for part in _INLINE.split(text):
        if not part:
            continue
        if part.startswith("**"):
            runs.append((part[2:-2], True, False))
        elif part.startswith("*"):
            runs.append((part[1:-1], False, True))
        else:
            runs.append((part, False, False))
    return runs


# --------------------------------------------------------------------------- #
# PDF rendering (reportlab)
# --------------------------------------------------------------------------- #


def _rl_markup(text: str) -> str:
    out = []
    for run, bold, italic in inline_runs(text):
        run = run.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
        if bold:
            run = f"<b>{run}</b>"
        if italic:
            run = f"<i>{run}</i>"
        out.append(run)
    return "".join(out).replace("\n", "<br/>")


def render_pdf(blocks: list[Block], target: Path) -> None:
    styles = getSampleStyleSheet()
    body = ParagraphStyle("body", parent=styles["Normal"], fontSize=10, leading=14)
    h = {
        1: ParagraphStyle("h1", parent=styles["Heading1"], fontSize=18, spaceAfter=6),
        2: ParagraphStyle("h2", parent=styles["Heading2"], fontSize=13, spaceBefore=10, spaceAfter=4),
        3: ParagraphStyle("h3", parent=styles["Heading3"], fontSize=11, spaceBefore=8, spaceAfter=2),
    }
    centered = ParagraphStyle("centered", parent=h[1], alignment=TA_CENTER)
    story = []
    first_heading = True
    for b in blocks:
        if b.kind == "heading":
            style = h.get(min(b.level, 3), h[3])
            if first_heading and b.level == 1:
                style = centered
            first_heading = False
            story.append(Paragraph(_rl_markup(b.text), style))
        elif b.kind == "paragraph":
            story.append(Paragraph(_rl_markup(b.text), body))
            story.append(Spacer(1, 4))
        elif b.kind in ("bullets", "numbers"):
            items = [ListItem(Paragraph(_rl_markup(t), body), leftIndent=12) for t in b.items]
            bullet = "bullet" if b.kind == "bullets" else "1"
            story.append(ListFlowable(items, bulletType=bullet, start=None if bullet == "bullet" else 1, leftIndent=14))
            story.append(Spacer(1, 4))
        elif b.kind == "table":
            data = [[Paragraph(_rl_markup(c), body) for c in row] for row in b.rows]
            table = Table(data, hAlign="LEFT")
            table.setStyle(
                TableStyle(
                    [
                        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                        ("BACKGROUND", (0, 0), (-1, 0), colors.whitesmoke),
                        ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ]
                )
            )
            story.append(table)
            story.append(Spacer(1, 6))
        elif b.kind == "rule":
            story.append(HRFlowable(width="100%", thickness=0.5, color=colors.grey, spaceBefore=6, spaceAfter=6))
    doc = SimpleDocTemplate(
        str(target),
        pagesize=A4,
        leftMargin=20 * mm,
        rightMargin=20 * mm,
        topMargin=18 * mm,
        bottomMargin=18 * mm,
        title=target.stem,
        author="Career AI Toolkit test kit (fictional)",
    )
    doc.build(story)


# --------------------------------------------------------------------------- #
# DOCX rendering (python-docx)
# --------------------------------------------------------------------------- #


def _docx_runs(paragraph, text: str) -> None:
    for run, bold, italic in inline_runs(text):
        r = paragraph.add_run(run)
        r.bold = bold
        r.italic = italic


def render_docx(blocks: list[Block], target: Path) -> None:
    doc = Document()
    doc.styles["Normal"].font.size = Pt(11)
    first_heading = True
    for b in blocks:
        if b.kind == "heading":
            p = doc.add_heading(level=min(b.level, 3))
            _docx_runs(p, b.text)
            if first_heading and b.level == 1:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            first_heading = False
        elif b.kind == "paragraph":
            p = doc.add_paragraph()
            for j, line in enumerate(b.text.split("\n")):
                if j:
                    p.add_run().add_break()
                _docx_runs(p, line)
        elif b.kind == "bullets":
            for t in b.items:
                _docx_runs(doc.add_paragraph(style="List Bullet"), t)
        elif b.kind == "numbers":
            for t in b.items:
                _docx_runs(doc.add_paragraph(style="List Number"), t)
        elif b.kind == "table":
            table = doc.add_table(rows=len(b.rows), cols=len(b.rows[0]))
            table.style = "Table Grid"
            for r, row in enumerate(b.rows):
                for c, cell in enumerate(row):
                    _docx_runs(table.cell(r, c).paragraphs[0], cell)
        elif b.kind == "rule":
            doc.add_paragraph("⸻").alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.save(str(target))


# --------------------------------------------------------------------------- #


def main() -> None:
    for rel, fmt in MANIFEST.items():
        source = SOURCES / rel
        if not source.is_file():
            raise SystemExit(f"Missing Markdown source: {source}")
        target = source.with_suffix(f".{fmt}")
        blocks = parse_markdown(source.read_text(encoding="utf-8"))
        if fmt == "pdf":
            render_pdf(blocks, target)
        elif fmt == "docx":
            render_docx(blocks, target)
        else:
            raise SystemExit(f"Unsupported format {fmt!r} for {rel}")
        print(f"{rel} -> {target.name}")


if __name__ == "__main__":
    main()
